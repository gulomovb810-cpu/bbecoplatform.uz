"""Grafiklar va jarayon sxemasi (PNG) — hujjat va taqdimot uchun."""
import json, sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from openpyxl import load_workbook

XLSX, OUT = sys.argv[1], sys.argv[2]
wb = load_workbook(XLSX, data_only=True)
pg = wb["Prognoz"]
rows = {}
for r in pg.iter_rows(min_row=4, values_only=True):
    if r[0]:
        rows[r[0]] = list(r[1:9])
years = [str(2027 + i) for i in range(8)]
S1, S2, S3 = "#2a78d6", "#1baf7a", "#eb6834"
INK, INK2, GRID, SURF = "#0b0b0b", "#52514e", "#e4e3df", "#ffffff"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11, "axes.edgecolor": GRID,
                     "axes.labelcolor": INK2, "xtick.color": INK2, "ytick.color": INK2})


def clean(ax):
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.grid(axis="y", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.tick_params(length=0)


# 1) Daromad va EBITDA
rev = [v / 1000 for v in rows["JAMI DAROMAD"]]
ebitda = [v / 1000 for v in rows["EBITDA"]]
ni = [v / 1000 for v in rows["SOF FOYDA"]]
fig, ax = plt.subplots(figsize=(10, 5), dpi=160)
x = range(len(years))
w = 0.38
ax.bar([i - w / 2 for i in x], rev, w - 0.03, color=S1, label="Daromad")
ax.bar([i + w / 2 for i in x], ebitda, w - 0.03, color=S2, label="EBITDA")
ax.axhline(0, color=INK2, linewidth=0.8)
for i in (1, 4):
    ax.text(i - w / 2, rev[i] + 50, f"{rev[i]:,.0f}", ha="center", color=INK, fontsize=10)
    ax.text(i + w / 2, ebitda[i] + 50, f"{ebitda[i]:,.0f}", ha="center", color=INK, fontsize=10)
ax.set_xticks(list(x), years)
ax.set_ylabel("ming $")
ax.set_title("Daromad va EBITDA, ming $", loc="left", color=INK, fontsize=13, fontweight="bold")
ax.legend(frameon=False, loc="upper left")
clean(ax)
fig.tight_layout()
fig.savefig(f"{OUT}/chart_revenue.png")
plt.close(fig)

# 2) Kumulyativ FCF
cf = [v / 1000 for v in rows["Kumulyativ FCF"]]
fig, ax = plt.subplots(figsize=(10, 4.6), dpi=160)
ax.plot(years, cf, color=S1, linewidth=2.2, marker="o", markersize=7, markeredgecolor=SURF, markeredgewidth=2)
ax.axhline(0, color=INK2, linewidth=0.8)
ax.fill_between(years, cf, 0, where=[v < 0 for v in cf], color=S3, alpha=0.10, interpolate=True)
ax.fill_between(years, cf, 0, where=[v >= 0 for v in cf], color=S2, alpha=0.12, interpolate=True)
ax.annotate(f"Eng chuqur nuqta: {min(cf):,.0f}", (cf.index(min(cf)), min(cf)), xytext=(0, -24),
            textcoords="offset points", ha="center", color=INK, fontsize=10)
ax.annotate(f"7-yil oxiri: +{cf[-1]:,.0f}", (7, cf[-1]), xytext=(-10, 10), textcoords="offset points",
            ha="right", color=INK, fontsize=10)
ax.set_ylabel("ming $")
ax.set_title("Kumulyativ erkin pul oqimi (loyiha), ming $", loc="left", color=INK, fontsize=13, fontweight="bold")
clean(ax)
ax.set_ylim(min(cf) * 1.3, cf[-1] * 1.25)
fig.tight_layout()
fig.savefig(f"{OUT}/chart_cashflow.png")
plt.close(fig)

# 3) Paket tannarxi tarkibi
tn = wb["Tannarx"]
items = []
start = False
for r in tn.iter_rows(min_row=4, values_only=True):
    if r[0] and str(r[0]).startswith("B)"):
        start = True
        continue
    if start and r[0] and r[0].startswith("Paket/plyonka o'zgaruvchan"):
        break
    if start and r[0] and isinstance(r[2], (int, float)) and r[2] > 1:
        items.append((r[0], r[2]))
items.sort(key=lambda t: t[1])
fig, ax = plt.subplots(figsize=(10, 4.4), dpi=160)
ax.barh([t[0] for t in items], [t[1] for t in items], color=S1, height=0.6)
for i, (_, v) in enumerate(items):
    ax.text(v + 15, i, f"${v:,.0f}", va="center", color=INK, fontsize=10)
ax.set_title("1 t kompost bo'ladigan paket tannarxi tarkibi, \\$ (sotish narxi \\$3 900)", loc="left",
             color=INK, fontsize=13, fontweight="bold")
for s in ("top", "right", "bottom"):
    ax.spines[s].set_visible(False)
ax.tick_params(length=0)
ax.set_xticks([])
fig.tight_layout()
fig.savefig(f"{OUT}/chart_unitcost.png")
plt.close(fig)

# 4) Texnologik jarayon sxemasi
steps = [
    ("1. Xomashyo", "Ishlatilgan yog',\nsoapstok, melassa\n→ filtrlash, sterilizatsiya"),
    ("2. Urug' ekish", "Shtamm (C. necator H16)\nkolba → 30 L → 1 m³\n30 °C, 24–30 soat"),
    ("3. Fermentatsiya", "10 m³ fed-batch\n30 °C, pH 6,8, DO >20%\n48–72 soat; N cheklanadi\n→ PHA hujayrada 70–80%"),
    ("4. Ajratish", "Separator → biomassa\nhujayrani parchalash\n(NaOH/SDS) → yuvish"),
    ("5. Quritish", "Purkab quritish\n→ PHA kukuni\n(tozaligi ≥95%)"),
    ("6. Kompaundlash", "PHA + PBAT + PLA\n+ qo'shimchalar\nikki shnekli ekstruder\n→ granula"),
    ("7. Mahsulot", "Plyonka puflash\n→ paket, qop,\nmulcha plyonka\n→ sifat nazorati"),
]
fig, ax = plt.subplots(figsize=(16, 4.0), dpi=160)
ax.set_xlim(0, 16)
ax.set_ylim(0, 4.0)
ax.axis("off")
bw, gap = 2.0, 0.27
colors = ["#e3eef3", "#e3eef3", "#d4ecdf", "#e3eef3", "#e3eef3", "#fbe5d9", "#fbe5d9"]
edges = ["#0b3d5c", "#0b3d5c", "#1f6f5c", "#0b3d5c", "#0b3d5c", "#b4501f", "#b4501f"]
for i, (h, b) in enumerate(steps):
    x0 = 0.15 + i * (bw + gap)
    box = FancyBboxPatch((x0, 0.55), bw, 2.75, boxstyle="round,pad=0.02,rounding_size=0.15",
                         facecolor=colors[i], edgecolor=edges[i], linewidth=1.6)
    ax.add_patch(box)
    ax.text(x0 + bw / 2, 2.95, h, ha="center", va="center", fontsize=12, fontweight="bold", color=edges[i])
    ax.text(x0 + bw / 2, 1.8, b, ha="center", va="center", fontsize=9.6, color=INK, linespacing=1.45)
    if i < len(steps) - 1:
        ax.add_patch(FancyArrowPatch((x0 + bw + 0.02, 1.9), (x0 + bw + gap - 0.02, 1.9),
                                     arrowstyle="-|>", mutation_scale=14, color=INK2, linewidth=1.4))
ax.text(0.15, 3.6, "Biotexnologik qism (PHA ishlab chiqarish)", fontsize=11, color="#1f6f5c", fontweight="bold")
ax.text(0.15 + 5 * (bw + gap), 3.6, "Polimer qayta ishlash (tayyor mahsulot)", fontsize=11, color="#b4501f",
        fontweight="bold")
ax.text(0.15, 0.15, "Oqova suv → dastlabki tozalash → kanalizatsiya;  qolgan biomassa → kompost/biogaz;  "
        "bug' kondensati → qozonga qaytariladi", fontsize=9.5, color=INK2)
fig.tight_layout()
fig.savefig(f"{OUT}/process_flow.png")
plt.close(fig)
print("ok")
