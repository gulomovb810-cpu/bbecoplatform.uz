"""BioPlast Tech moliyaviy modeli (openpyxl, formulalar bilan)."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
RESERVE = [int(x) for x in sys.argv[2].split(",")]
wb = Workbook()

F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
H1 = Font(name=F, bold=True, size=14, color="FFFFFF")
HDR = Font(name=F, bold=True, color="FFFFFF")
FILL_H1 = PatternFill("solid", fgColor="0B3D5C")
FILL_HDR = PatternFill("solid", fgColor="1F6F5C")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_SEC = PatternFill("solid", fgColor="E3EEF3")
FILL_TOT = PatternFill("solid", fgColor="D9EAD3")
thin = Side(style="thin", color="B7C4CC")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
WRAP = Alignment(wrap_text=True, vertical="top")


def title(ws, text, width_cols):
    ws["A1"] = text
    ws["A1"].font = H1
    for c in range(1, width_cols + 1):
        ws.cell(row=1, column=c).fill = FILL_H1
    ws.row_dimensions[1].height = 24


def header(ws, row, labels):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX


# ---------------------------------------------------------------- Kirish
ki = wb.active
ki.title = "Kirish"
title(ki, "BioPlast Tech — asosiy farazlar (ko'k = kiritiladigan qiymat, sariq = asosiy faraz)", 4)
ki["A2"] = ("Narxlar AQSh dollarida, QQSsiz, real (inflyatsiyasiz) ifodada. "
            "Ko'k raqamlarni o'zgartirsangiz, barcha varaqlar avtomatik qayta hisoblanadi.")
ki["A2"].font = Font(name=F, italic=True)
ki.merge_cells("A2:D2")
header(ki, 3, ["Ko'rsatkich", "Qiymat", "Birlik", "Manba / izoh"])

K = {}  # key -> absolute ref
rows = [
    ("sec", "Makroiqtisodiy"),
    ("fx", "Valyuta kursi", 12000, "so'm / $", "O'zR MB: 22.10.2026 da 11 966 so'm; yaxlitlangan. cbu.uz"),
    ("el_uzs", "Elektr energiya tarifi (yuridik shaxslar)", 1100, "so'm / kWh", "VM qarori, 01.06.2026 dan (kun.uz, 15.05.2026; interfax.com)"),
    ("gas_uzs", "Tabiiy gaz tarifi (yuridik shaxslar)", 2000, "so'm / m³", "01.06.2026 dan (interfax.com, 15.05.2026)"),
    ("tax", "Foyda solig'i", 0.15, "%", "O'zR Soliq kodeksi, 337-modda — 15%"),
    ("soc", "Ijtimoiy soliq", 0.12, "%", "O'zR Soliq kodeksi — 12% (ish haqi fondiga)"),
    ("disc", "Diskont stavkasi (USD, xavf bilan)", 0.15, "%", "Startap xavfi hisobga olingan faraz"),
    ("sec", "PHA ishlab chiqarish (Cupriavidus necator H16)"),
    ("yield", "PHA chiqishi (substrat bo'yicha)", 0.55, "kg PHA / kg yog'", "Nazariy maks. ~0.7; o'simlik yog'ida 0.6–0.8 qayd etilgan, konservativ"),
    ("oil", "Xomashyo: ishlatilgan yog' / soapstok narxi", 500, "$ / t", "Faraz: yog'-moy kombinatlari va HoReCa'dan kelishuv narxi — KP bilan tasdiqlash kerak"),
    ("nutr", "Ozuqa tuzlari (N, P, mikroelementlar)", 120, "$ / t PHA", "Faraz (ammoniy sulfat, KH2PO4, MgSO4, iz elementlar)"),
    ("dsp", "Ajratish kimyoviy moddalari (NaOH, SDS/NaClO)", 300, "$ / t PHA", "Faraz; erituvchisiz ajratish usuli"),
    ("kwh_pha", "Elektr sarfi", 9000, "kWh / t PHA", "Aeratsiya, aralashtirish, separator, quritish — kichik masshtab"),
    ("gas_pha", "Gaz sarfi (bug' qozoni)", 1900, "m³ / t PHA", "~25 t bug' / t PHA, 75 m³ gaz / t bug'"),
    ("water_pha", "Suv va oqova suv", 60, "$ / t PHA", "Faraz"),
    ("pack_pha", "Qadoqlash (25 kg qop)", 30, "$ / t PHA", "Faraz"),
    ("sec", "Kompost bo'ladigan plyonka / paket (PBAT + PLA + PHA)"),
    ("pbat", "PBAT narxi (O'zbekistonga yetkazilgan)", 1700, "$ / t", "Xitoy: 9 850 yuan/t (CICC, 21.09.2025) + logistika va bojxona"),
    ("pla", "PLA narxi (yetkazilgan)", 2600, "$ / t", "Xitoy: 2 010–2 520 $/t (Intratec, 2025) + logistika"),
    ("add", "Qo'shimchalar (CaCO3, kompatibilizator, masterbatch)", 900, "$ / t", "Faraz"),
    ("add_sh", "Qo'shimchalar ulushi", 0.08, "%", "Retsept"),
    ("pbat_sh", "PBAT : (PBAT+PLA) nisbati", 0.72, "%", "Retsept — qolgani PLA"),
    ("scrap", "Ishlab chiqarish chiqindisi (yo'qotish)", 0.05, "%", "Qisman qayta ishlatiladi"),
    ("kwh_film", "Elektr sarfi (kompaund + plyonka + paket)", 900, "kWh / t", "Faraz"),
    ("pack_film", "Qadoqlash, bo'yoq", 110, "$ / t", "Faraz"),
    ("sec", "Sotish narxlari (QQSsiz, zavoddan)"),
    ("p_film", "Kompost bo'ladigan paket/plyonka narxi", 3900, "$ / t", "Polietilen paketdan ~2 baravar qimmat; import analoglaridan arzon"),
    ("p_pha", "Toza PHA granula narxi", 4500, "$ / t", "IMARC: 3 443–3 712 $/t (12.2025); Mordor: 4 000–15 000 $/t"),
    ("sell", "Sotish va logistika xarajati", 0.04, "% daromaddan", "Faraz"),
    ("sec", "Doimiy xarajatlar"),
    ("rent_m2", "Ishlab chiqarish maydoni", 1500, "m²", "Toshkent viloyati sanoat zonasi"),
    ("rent", "Ijara narxi", 4, "$ / m² / oy", "Faraz — mahalliy taklif bilan tasdiqlash"),
    ("maint", "Ta'mirlash va ehtiyot qismlar", 0.03, "% uskuna qiymatidan / yil", "Faraz"),
    ("mkt", "Marketing (doimiy qism)", 30000, "$ / yil", "Ko'rgazmalar, sertifikat markirovkasi, raqamli marketing"),
    ("admin", "Ma'muriy, sug'urta, IT", 20000, "$ / yil", "Faraz"),
    ("cert", "Sertifikatni yangilash, laboratoriya sinovlari", 10000, "$ / yil", "TÜV AUSTRIA / DIN CERTCO audit"),
    ("rnd", "Ilmiy-tadqiqot (shtamm, retsept)", 25000, "$ / yil", "TKTI bilan hamkorlik"),
    ("util", "Kommunal (yoritish, isitish)", 12000, "$ / yil", "Faraz"),
    ("sec", "Kapital va moliyalashtirish"),
    ("life", "Uskunalarning foydali xizmat muddati", 10, "yil", "To'g'ri chiziqli amortizatsiya"),
    ("wc", "Aylanma kapital", 0.15, "% daromad o'sishidan", "Zaxira + debitorlar - kreditorlar"),
    ("loan_sh", "Bank krediti ulushi (CAPEXdan)", 0.5, "%", "Masalan: Biznesni rivojlantirish banki, Sanoatqurilishbank"),
    ("rate", "Kredit foizi (USD)", 0.10, "% / yil", "Faraz; aniq stavka bankdan olinadi"),
    ("grant", "Grant (0-yil)", 50000, "$", "Innovatsion grant — tasdiqlanmagan, faraz"),
]
r = 4
for item in rows:
    if item[0] == "sec":
        c = ki.cell(row=r, column=1, value=item[1])
        c.font = BOLD
        for col in range(1, 5):
            ki.cell(row=r, column=col).fill = FILL_SEC
        r += 1
        continue
    key, label, val, unit, note = item
    ki.cell(row=r, column=1, value=label).font = BLACK
    v = ki.cell(row=r, column=2, value=val)
    v.font = BLUE
    if unit.startswith("%"):
        v.number_format = PCT
    elif "$" in unit and val < 100 and isinstance(val, (int, float)) and key == "rent":
        v.number_format = USD2
    elif "$" in unit:
        v.number_format = USD
    else:
        v.number_format = NUM
    ki.cell(row=r, column=3, value=unit).font = BLACK
    ki.cell(row=r, column=4, value=note).font = Font(name=F, size=9, color="555555")
    ki.cell(row=r, column=4).alignment = WRAP
    if key in ("p_film", "oil", "p_pha", "yield", "pbat"):
        v.fill = FILL_KEY
    K[key] = f"Kirish!$B${r}"
    r += 1
# Hosilaviy
ki.cell(row=r, column=1, value="Hosilaviy qiymatlar").font = BOLD
for col in range(1, 5):
    ki.cell(row=r, column=col).fill = FILL_SEC
r += 1
for key, label, f, fmt, unit in [
    ("el_usd", "Elektr narxi", f"={K['el_uzs']}/{K['fx']}", '$0.0000', "$ / kWh"),
    ("gas_usd", "Gaz narxi", f"={K['gas_uzs']}/{K['fx']}", '$0.0000', "$ / m³"),
]:
    ki.cell(row=r, column=1, value=label).font = BLACK
    c = ki.cell(row=r, column=2, value=f)
    c.font = BLACK
    c.number_format = fmt
    ki.cell(row=r, column=3, value=unit)
    K[key] = f"Kirish!$B${r}"
    r += 1
ki.column_dimensions["A"].width = 52
ki.column_dimensions["B"].width = 14
ki.column_dimensions["C"].width = 24
ki.column_dimensions["D"].width = 70
ki.freeze_panes = "A4"

# ---------------------------------------------------------------- CAPEX
cx = wb.create_sheet("CAPEX")
title(cx, "Kapital xarajatlar (CAPEX) — uskunalar va yetkazib beruvchilar", 8)
cx["A2"] = ("Narxlar — 2025–2026 yillardagi ochiq bozor e'lonlari va ishlab chiqaruvchi narx diapazonlari asosidagi taxmin. "
            "Sotib olishdan oldin kamida 3 ta tijorat taklifi (KP) olinishi shart.")
cx["A2"].font = Font(name=F, italic=True, size=9)
cx.merge_cells("A2:H2")
header(cx, 3, ["№", "Uskuna / xarajat", "Faza", "Turi (U=uskuna, B=boshqa)", "Soni", "Narx ($)", "Jami ($)", "Yetkazib beruvchilar (misol)"])
capex_items = [
    # Faza 1 — laboratoriya, pilot, plyonka liniyasi
    ("Stol usti bioreaktor 5 L (pH, DO, harorat nazorati)", 1, "U", 2, 12000, "Sartorius BIOSTAT A; Eppendorf BioFlo 120; Bailun BLBIO-5GJ"),
    ("Avtoklav 75 L", 1, "U", 1, 7000, "Tuttnauer; Systec; Xitoy: Zealway"),
    ("Chayqatuvchi inkubator (shaker)", 1, "U", 1, 6000, "Eppendorf Innova; Infors HT; Zhicheng"),
    ("Laminar boks (II klass)", 1, "U", 1, 5000, "Esco; Thermo Fisher; Biobase"),
    ("Laboratoriya sentrifugasi", 1, "U", 1, 5000, "Eppendorf; Hettich; Hunan Xiangyi"),
    ("Mikroskop (fazali kontrast)", 1, "U", 1, 2500, "Olympus/Evident; Zeiss; Optika"),
    ("Spektrofotometr UV-Vis (OD600)", 1, "U", 1, 4000, "Shimadzu; Thermo; Biobase"),
    ("Gaz xromatograf GC-FID (PHA miqdori)", 1, "U", 1, 35000, "Agilent 8860; Shimadzu GC-2030; Fuli"),
    ("Quritish shkafi, analitik tarozi, pH-metr", 1, "U", 1, 5000, "Memmert; Ohaus; Mettler Toledo"),
    ("Erish oqimi indeksi (MFI) asbobi", 1, "U", 1, 4000, "Instron/CEAST; Xitoy: Chengde"),
    ("Universal cho'zilish sinov mashinasi", 1, "U", 1, 12000, "Instron; Zwick; WANCE"),
    ("Pilot fermenter 300 L + urug' 30 L", 1, "U", 1, 45000, "Bailun (BLBIO); T&J Bio-engineering; Zhengzhou Laboao"),
    ("Pilot ajratish: quvurli sentrifuga + vakuum quritgich", 1, "U", 1, 22000, "Liaoyang Shenzhou; Labfreez"),
    ("Ikki shnekli kompaundlash ekstruderi 65 mm (150–250 kg/soat) + granulyator", 1, "U", 1, 85000, "Coperion ZSK; Nanjing Kerke; Nanjing Haisi"),
    ("Biodegradable plyonka puflash liniyasi №1 (65 mm, ABA, ~90 kg/soat)", 1, "U", 1, 70000, "Ruian Ruipai; Wenzhou Ruida; Hosokawa Alpine (Yevropa)"),
    ("Paket tayyorlash mashinasi (mayka-paket, rulondagi qop)", 1, "U", 2, 18000, "Ruian Xinye; Ruian Kingdom"),
    ("Fleksografik bosma mashina (2 rang)", 1, "U", 1, 22000, "Ruian Changhong; Ruian Ruiting"),
    ("Elektr tarmog'i va siqilgan havo (1-faza)", 1, "U", 1, 20000, "Mahalliy pudratchi"),
    ("Binoni ta'mirlash, toza xona, ventilyatsiya", 1, "B", 1, 80000, "Mahalliy qurilish pudratchisi"),
    ("Xalqaro sertifikatlash (OK compost INDUSTRIAL/HOME, OK biodegradable SOIL)", 1, "B", 1, 25000, "TÜV AUSTRIA; DIN CERTCO"),
    ("Milliy sertifikat, sanitariya-gigiyena xulosasi", 1, "B", 1, 5000, "O'zstandart agentligi; SES"),
    ("Loyihalash, ruxsatnomalar, ekologik ekspertiza", 1, "B", 1, 30000, "Loyiha instituti; Ekologiya vazirligi"),
    # Faza 2 — namoyish zavodi 2×10 m3
    ("Urug' fermenteri 1 m³", 2, "U", 1, 40000, "Bailun; T&J; Laboao"),
    ("Asosiy fermenter 10 m³ (SIP/CIP, avtomatika)", 2, "U", 2, 90000, "Bailun; T&J; Jiangsu Zhongtai; Yevropa: Bioengineering AG"),
    ("Muhit tayyorlash, oziqlantirish idishlari, sterilizator", 2, "U", 1, 30000, "Wenzhou Ace; Mahalliy"),
    ("Havo puflagich + steril filtrlar", 2, "U", 1, 35000, "Atlas Copco ZS; Kaeser; Pall/Donaldson filtrlari"),
    ("Bug' qozoni 1 t/soat (gaz)", 2, "U", 1, 35000, "Bosch Industriekessel; Xitoy: ZOZEN; Mahalliy"),
    ("Sovutish minorasi + chiller", 2, "U", 1, 30000, "Mahalliy; Xitoy"),
    ("Suv tayyorlash (teskari osmos)", 2, "U", 1, 15000, "Mahalliy; Xitoy"),
    ("Disk separatori (biomassa ajratish)", 2, "U", 1, 70000, "Alfa Laval; GEA; Xitoy: Nanjing Huasheng"),
    ("Hujayrani parchalash/yuvish reaktorlari", 2, "U", 1, 30000, "Wenzhou Ace; Mahalliy"),
    ("Yuvish uchun 2-separator", 2, "U", 1, 40000, "Alfa Laval; GEA; Nanjing Huasheng"),
    ("Purkab quritgich (50 kg/soat bug'lanish)", 2, "U", 1, 60000, "GEA Niro; Changzhou Lemar; Changzhou Yibu"),
    ("Biodegradable plyonka puflash liniyasi №2 (65 mm, ~90 kg/soat)", 2, "U", 1, 70000, "Ruian Ruipai; Wenzhou Ruida"),
    ("Oqova suvni dastlabki tozalash", 2, "U", 1, 40000, "Mahalliy; Xitoy"),
    ("PLC/SCADA avtomatlashtirish", 2, "U", 1, 25000, "Siemens; Schneider Electric; Delta"),
]
r0 = 4
for i, (name, ph, typ, qty, price, sup) in enumerate(capex_items, 1):
    r = r0 + i - 1
    cx.cell(row=r, column=1, value=i).font = BLACK
    cx.cell(row=r, column=2, value=name).font = BLACK
    cx.cell(row=r, column=3, value=ph).font = BLUE
    cx.cell(row=r, column=4, value=typ).font = BLUE
    cx.cell(row=r, column=5, value=qty).font = BLUE
    p = cx.cell(row=r, column=6, value=price)
    p.font = BLUE
    p.number_format = USD
    t = cx.cell(row=r, column=7, value=f"=E{r}*F{r}")
    t.font = BLACK
    t.number_format = USD
    cx.cell(row=r, column=8, value=sup).font = Font(name=F, size=9)
rl = r0 + len(capex_items) - 1
rng = lambda col: f"${col}${r0}:${col}${rl}"
r = rl + 2
cx.cell(row=r, column=2, value="Jamlanma").font = BOLD
cx.cell(row=r, column=6, value="1-faza ($)").font = BOLD
cx.cell(row=r, column=7, value="2-faza ($)").font = BOLD
cx.cell(row=r, column=8, value="Jami ($)").font = BOLD
r += 1
CX = {}
sumrows = [
    ("eq", "Uskunalar", lambda ph: f"=SUMIFS({rng('G')},{rng('C')},{ph},{rng('D')},\"U\")"),
    ("frt", "Yetkazish, bojxona rasmiylashtiruvi", None),
    ("inst", "Montaj, quvurlar, elektr ulanish", None),
    ("oth", "Boshqa (bino, sertifikat, loyihalash)", lambda ph: f"=SUMIFS({rng('G')},{rng('C')},{ph},{rng('D')},\"B\")"),
    ("sub", "Oraliq jami", None),
    ("cont", "Kutilmagan xarajatlar", None),
    ("tot", "JAMI CAPEX", None),
]
# rates rows (inputs) placed below
rate_row = r + len(sumrows) + 1
cx.cell(row=rate_row, column=2, value="Yetkazish va bojxona stavkasi (uskunadan)").font = BLACK
cx.cell(row=rate_row, column=6, value=0.12).font = BLUE
cx.cell(row=rate_row, column=6).number_format = PCT
cx.cell(row=rate_row, column=8, value="Texnologik uskunalar uchun boj imtiyozlari bo'lishi mumkin — konservativ hisob").font = Font(name=F, size=9)
cx.cell(row=rate_row + 1, column=2, value="Montaj stavkasi (uskunadan)").font = BLACK
cx.cell(row=rate_row + 1, column=6, value=0.15).font = BLUE
cx.cell(row=rate_row + 1, column=6).number_format = PCT
cx.cell(row=rate_row + 2, column=2, value="Kutilmagan xarajatlar stavkasi").font = BLACK
cx.cell(row=rate_row + 2, column=6, value=0.10).font = BLUE
cx.cell(row=rate_row + 2, column=6).number_format = PCT
FRT, INST, CONT = f"$F${rate_row}", f"$F${rate_row+1}", f"$F${rate_row+2}"
for j, (key, label, fn) in enumerate(sumrows):
    rr = r + j
    CX[key] = rr
    cx.cell(row=rr, column=2, value=label).font = BOLD if key in ("sub", "tot") else BLACK
    for col, ph in ((6, 1), (7, 2)):
        cl = L(col)
        if fn:
            f = fn(ph)
        elif key == "frt":
            f = f"={cl}{CX['eq']}*{FRT}"
        elif key == "inst":
            f = f"={cl}{CX['eq']}*{INST}"
        elif key == "sub":
            f = f"=SUM({cl}{CX['eq']}:{cl}{CX['oth']})"
        elif key == "cont":
            f = f"={cl}{CX['sub']}*{CONT}"
        elif key == "tot":
            f = f"={cl}{CX['sub']}+{cl}{CX['cont']}"
        c = cx.cell(row=rr, column=col, value=f)
        c.number_format = USD
        c.font = BOLD if key in ("sub", "tot") else BLACK
    c = cx.cell(row=rr, column=8, value=f"=F{rr}+G{rr}")
    c.number_format = USD
    c.font = BOLD if key in ("sub", "tot") else BLACK
    if key == "tot":
        for col in range(2, 9):
            cx.cell(row=rr, column=col).fill = FILL_TOT
# amortizatsiya qilinadigan baza = jami CAPEX (sertifikat va loyihalash ham kapitallashtiriladi — soddalashtirish)
CAPEX1 = f"CAPEX!$F${CX['tot']}"
CAPEX2 = f"CAPEX!$G${CX['tot']}"
CAPEXT = f"CAPEX!$H${CX['tot']}"
EQ1 = f"CAPEX!$F${CX['eq']}"
EQ2 = f"CAPEX!$G${CX['eq']}"
for col, w in zip("ABCDEFGH", (5, 62, 7, 12, 7, 12, 13, 58)):
    cx.column_dimensions[col].width = w
cx.freeze_panes = "A4"

# ---------------------------------------------------------------- Xodimlar
xo = wb.create_sheet("Xodimlar")
title(xo, "Xodimlar va ish haqi fondi", 7)
header(xo, 3, ["Lavozim", "1-yil soni", "2-yildan soni", "Oylik ish haqi ($, brutto)", "1-yil fondi ($)", "2-yildan fondi ($)", "Izoh"])
staff = [
    ("Direktor", 1, 1, 1500, ""),
    ("Bosh texnolog-mikrobiolog", 1, 1, 1200, "Biotexnologiya magistri; TKTI bitiruvchisi"),
    ("Mikrobiolog / laborant", 2, 3, 700, "Shtamm saqlash, sifat tahlili, GC"),
    ("Fermentatsiya va ajratish operatori", 2, 8, 550, "24/7, 4 smena × 2 kishi"),
    ("Ekstruziya va plyonka operatori", 6, 9, 550, "3 smena"),
    ("Paket kesish va qadoqlash ishchisi", 3, 6, 450, ""),
    ("Mexanik-elektrik (KIP)", 1, 2, 600, ""),
    ("Sifat nazorati muhandisi", 1, 1, 650, "ISO 9001 tizimi"),
    ("Savdo va marketing menejeri", 1, 2, 700, "B2B savdo, eksport"),
    ("Bosh hisobchi", 1, 1, 700, ""),
    ("Omborchi-logist", 1, 1, 450, ""),
]
for i, (pos, n1, n2, sal, note) in enumerate(staff):
    rr = 4 + i
    xo.cell(row=rr, column=1, value=pos).font = BLACK
    for col, v in ((2, n1), (3, n2), (4, sal)):
        c = xo.cell(row=rr, column=col, value=v)
        c.font = BLUE
        c.number_format = USD if col == 4 else NUM
    xo.cell(row=rr, column=5, value=f"=B{rr}*D{rr}*12").number_format = USD
    xo.cell(row=rr, column=6, value=f"=C{rr}*D{rr}*12").number_format = USD
    xo.cell(row=rr, column=7, value=note).font = Font(name=F, size=9)
sl = 4 + len(staff) - 1
rr = sl + 1
xo.cell(row=rr, column=1, value="Jami").font = BOLD
for col in "BCEF":
    c = xo[f"{col}{rr}"]
    c.value = f"=SUM({col}4:{col}{sl})"
    c.font = BOLD
    c.number_format = USD if col in "EF" else NUM
rr += 1
xo.cell(row=rr, column=1, value="Ijtimoiy soliq bilan jami fond").font = BOLD
for col in "EF":
    c = xo[f"{col}{rr}"]
    c.value = f"={col}{rr-1}*(1+{K['soc']})"
    c.font = BOLD
    c.number_format = USD
    c.fill = FILL_TOT
PAY1 = f"Xodimlar!$E${rr}"
PAY2 = f"Xodimlar!$F${rr}"
xo.cell(row=rr + 2, column=1, value="Ma'lumot uchun: O'zbekistonda o'rtacha oylik ish haqi 2026-yil I chorakda 6,8 mln so'm (~$570) — stat.uz / kun.uz.").font = Font(name=F, italic=True, size=9)
for col, w in zip("ABCDEFG", (40, 11, 13, 20, 16, 17, 40)):
    xo.column_dimensions[col].width = w

# ---------------------------------------------------------------- Tannarx
tn = wb.create_sheet("Tannarx")
title(tn, "1 tonna mahsulotning o'zgaruvchan tannarxi", 4)
header(tn, 3, ["Xarajat moddasi", "Hisob", "$ / t", "Ulush"])
tn["A4"] = "A) Toza PHA granula (1 t)"
tn["A4"].font = BOLD
for c in range(1, 5):
    tn.cell(row=4, column=c).fill = FILL_SEC
pha_lines = [
    ("Yog' xomashyosi", "1/chiqish × narx", f"=1/{K['yield']}*{K['oil']}"),
    ("Ozuqa tuzlari", "", f"={K['nutr']}"),
    ("Ajratish kimyoviy moddalari", "", f"={K['dsp']}"),
    ("Elektr energiya", "kWh × tarif", f"={K['kwh_pha']}*{K['el_usd']}"),
    ("Tabiiy gaz (bug')", "m³ × tarif", f"={K['gas_pha']}*{K['gas_usd']}"),
    ("Suv va oqova", "", f"={K['water_pha']}"),
    ("Qadoqlash", "", f"={K['pack_pha']}"),
]
for i, (lab, how, f) in enumerate(pha_lines):
    rr = 5 + i
    tn.cell(row=rr, column=1, value=lab)
    tn.cell(row=rr, column=2, value=how).font = Font(name=F, size=9, color="555555")
    tn.cell(row=rr, column=3, value=f).number_format = USD
    tn.cell(row=rr, column=4, value=f"=C{rr}/$C${5+len(pha_lines)}").number_format = PCT
pt = 5 + len(pha_lines)
tn.cell(row=pt, column=1, value="PHA o'zgaruvchan tannarxi").font = BOLD
c = tn.cell(row=pt, column=3, value=f"=SUM(C5:C{pt-1})")
c.font = BOLD
c.number_format = USD
c.fill = FILL_TOT
PHA_VC = f"Tannarx!$C${pt}"

r = pt + 2
tn.cell(row=r, column=1, value="B) Kompost bo'ladigan paket/plyonka (1 t, PHA ulushi 6,5% bo'lganda)").font = BOLD
for c in range(1, 5):
    tn.cell(row=r, column=c).fill = FILL_SEC
tn.cell(row=r + 1, column=1, value="Namuna retseptdagi PHA ulushi")
c = tn.cell(row=r + 1, column=3, value=0.065)
c.font = BLUE
c.number_format = PCT
SH = f"$C${r+1}"
base = f"(1-{SH}-{K['add_sh']})"
film_lines = [
    ("PBAT", "ulush × narx × (1+yo'qotish)", f"={base}*{K['pbat_sh']}*{K['pbat']}*(1+{K['scrap']})"),
    ("PLA", "", f"={base}*(1-{K['pbat_sh']})*{K['pla']}*(1+{K['scrap']})"),
    ("Qo'shimchalar", "", f"={K['add_sh']}*{K['add']}*(1+{K['scrap']})"),
    ("PHA (ichki tannarxda)", "", f"={SH}*{PHA_VC}*(1+{K['scrap']})"),
    ("Elektr energiya", "", f"={K['kwh_film']}*{K['el_usd']}"),
    ("Qadoqlash, bo'yoq", "", f"={K['pack_film']}"),
]
fs = r + 2
for i, (lab, how, f) in enumerate(film_lines):
    rr = fs + i
    tn.cell(row=rr, column=1, value=lab)
    tn.cell(row=rr, column=2, value=how).font = Font(name=F, size=9, color="555555")
    tn.cell(row=rr, column=3, value=f).number_format = USD
    tn.cell(row=rr, column=4, value=f"=C{rr}/$C${fs+len(film_lines)}").number_format = PCT
ft = fs + len(film_lines)
tn.cell(row=ft, column=1, value="Paket/plyonka o'zgaruvchan tannarxi").font = BOLD
c = tn.cell(row=ft, column=3, value=f"=SUM(C{fs}:C{ft-1})")
c.font = BOLD
c.number_format = USD
c.fill = FILL_TOT
tn.cell(row=ft + 1, column=1, value="Sotish narxi")
tn.cell(row=ft + 1, column=3, value=f"={K['p_film']}").number_format = USD
tn.cell(row=ft + 2, column=1, value="Marjinal foyda (sotish xarajatidan oldin)").font = BOLD
c = tn.cell(row=ft + 2, column=3, value=f"=C{ft+1}-C{ft}")
c.font = BOLD
c.number_format = USD
tn.cell(row=ft + 2, column=4, value=f"=C{ft+2}/C{ft+1}").number_format = PCT
# Prognoz uchun: PHA'siz qism (1 t plyonka uchun) — ulushga bog'liq formula Prognoz'da hisoblanadi
for col, w in zip("ABCD", (58, 30, 14, 10)):
    tn.column_dimensions[col].width = w

# ---------------------------------------------------------------- Prognoz
pg = wb.create_sheet("Prognoz")
NY = 8  # 0..7
years = [2027 + i for i in range(NY)]
title(pg, "Daromad, xarajat, foyda va pul oqimi prognozi ($)", NY + 2)
header(pg, 3, ["Ko'rsatkich"] + [f"{y}\n({i}-yil)" for i, y in enumerate(years)] + ["Izoh"])
pg.row_dimensions[3].height = 30
cols = [L(2 + i) for i in range(NY)]
NOTE = L(2 + NY)
row = [4]
R = {}


def sec(label):
    rr = row[0]
    pg.cell(row=rr, column=1, value=label).font = BOLD
    for c in range(1, NY + 3):
        pg.cell(row=rr, column=c).fill = FILL_SEC
    row[0] += 1


def line(key, label, vals, fmt=USD, bold=False, note="", inp=False, total=False):
    rr = row[0]
    R[key] = rr
    pg.cell(row=rr, column=1, value=label).font = BOLD if bold else BLACK
    for i, col in enumerate(cols):
        v = vals(i, col) if callable(vals) else vals[i]
        c = pg.cell(row=rr, column=2 + i, value=v)
        c.number_format = fmt
        c.font = BLUE if inp else (BOLD if bold else BLACK)
        if total:
            c.fill = FILL_TOT
    if note:
        pg.cell(row=rr, column=NY + 2, value=note).font = Font(name=F, size=9, color="555555")
    row[0] += 1


def prev(col):
    return L(pg[col + "1"].column - 1)


sec("Ishlab chiqarish hajmi")
line("film", "Paket/plyonka ishlab chiqarish (t)", [0, 300, 600, 800, 850, 850, 850, 850], NUM, inp=True,
     note="2 ta liniya × 90 kg/soat × 6 000 soat = 1 080 t maks. quvvat; barqaror yilda 79% yuklanish")
line("pha", "PHA ishlab chiqarish (t)", [0, 3, 25, 55, 58, 58, 58, 58], NUM1, inp=True,
     note="2027–28: 300 L pilot; 2029 o'rtasidan 2×10 m³ (≈160 partiya/yil × 0,36 t)")
line("share", "PHA ulushi aralashmada", [0, 0.01, 0.04, 0.065, 0.065, 0.065, 0.065, 0.065], PCT, inp=True,
     note="Standart paketda 6,5% PHA; premium (uyda kompost) liniyada ulush yuqori")
line("pha_use", "PHA ichki ishlatish (t)", lambda i, c: f"={c}{R['film']}*{c}{R['share']}", NUM1)
line("pha_sale", "PHA sotish (t)", lambda i, c: f"=MAX(0,{c}{R['pha']}-{c}{R['pha_use']})", NUM1,
     note="Ichki ehtiyojdan ortiqcha qism sotiladi")
line("pha_buy", "PHA yetishmovchiligi — sotib olinadi (t)", lambda i, c: f"=MAX(0,{c}{R['pha_use']}-{c}{R['pha']})", NUM1)

sec("Daromad")
line("rev_film", "Paket/plyonka sotuvi", lambda i, c: f"={c}{R['film']}*{K['p_film']}")
line("rev_pha", "Toza PHA sotuvi", lambda i, c: f"={c}{R['pha_sale']}*{K['p_pha']}")
line("rev", "JAMI DAROMAD", lambda i, c: f"={c}{R['rev_film']}+{c}{R['rev_pha']}", bold=True, total=True)

sec("O'zgaruvchan xarajatlar")
nonpha = (lambda c: f"((1-{c}{R['share']}-{K['add_sh']})*({K['pbat_sh']}*{K['pbat']}+(1-{K['pbat_sh']})*{K['pla']})"
          f"+{K['add_sh']}*{K['add']})*(1+{K['scrap']})+{K['kwh_film']}*{K['el_usd']}+{K['pack_film']}")
line("vc_film", "Plyonka xomashyosi va energiya (PHA'siz)", lambda i, c: f"={c}{R['film']}*({nonpha(c)})")
line("vc_pha", "PHA ishlab chiqarish xarajati", lambda i, c: f"={c}{R['pha']}*{PHA_VC}")
line("vc_buy", "Sotib olingan PHA", lambda i, c: f"={c}{R['pha_buy']}*{K['p_pha']}")
line("vc_sell", "Sotish va logistika", lambda i, c: f"={c}{R['rev']}*{K['sell']}")
line("vc", "Jami o'zgaruvchan xarajat", lambda i, c: f"=SUM({c}{R['vc_film']}:{c}{R['vc_sell']})", bold=True)
line("gp", "MARJINAL FOYDA", lambda i, c: f"={c}{R['rev']}-{c}{R['vc']}", bold=True, total=True)

sec("Doimiy xarajatlar")
line("k", "Faoliyat koeffitsienti", [0.5, 1, 1, 1, 1, 1, 1, 1], '0.00', inp=True,
     note="0-yil: qurilish va ilmiy ish davri — xarajatlar 50%")
line("fx_pay", "Ish haqi + ijtimoiy soliq",
     lambda i, c: f"={c}{R['k']}*IF({i}<=1,{PAY1},{PAY2})")
line("fx_rent", "Ijara", lambda i, c: f"={K['rent_m2']}*{K['rent']}*12")
line("fx_maint", "Ta'mirlash", lambda i, c: f"=IF({i}>=1,{EQ1},0)*{K['maint']}+IF({i}>=2,{EQ2},0)*{K['maint']}")
line("fx_oth", "Marketing, ma'muriy, sertifikat, ITI, kommunal",
     lambda i, c: f"={c}{R['k']}*({K['mkt']}+{K['admin']}+{K['cert']}+{K['rnd']}+{K['util']})")
line("fc", "Jami doimiy xarajat", lambda i, c: f"=SUM({c}{R['fx_pay']}:{c}{R['fx_oth']})", bold=True)

sec("Foyda")
line("ebitda", "EBITDA", lambda i, c: f"={c}{R['gp']}-{c}{R['fc']}", bold=True, total=True)
line("ebitda_m", "EBITDA rentabelligi", lambda i, c: f"=IF({c}{R['rev']}=0,0,{c}{R['ebitda']}/{c}{R['rev']})", PCT)
line("dep", "Amortizatsiya",
     lambda i, c: f"=IF({i}>=1,{CAPEX1},0)/{K['life']}+IF({i}>=2,{CAPEX2},0)/{K['life']}")
line("ebit", "Foiz va soliqqacha foyda (EBIT)", lambda i, c: f"={c}{R['ebitda']}-{c}{R['dep']}")
# financing rows referenced later; placeholders filled after
int_row_placeholder = row[0]
line("int", "Kredit foizi", [0] * NY)
line("ebt", "Soliqqacha foyda", lambda i, c: f"={c}{R['ebit']}-{c}{R['int']}")
line("taxp", "Foyda solig'i", lambda i, c: f"=MAX(0,{c}{R['ebt']})*{K['tax']}",
     note="Soddalashtirish: zarar keyingi yillarga o'tkazilmaydi (konservativ)")
line("ni", "SOF FOYDA", lambda i, c: f"={c}{R['ebt']}-{c}{R['taxp']}", bold=True, total=True)

sec("Loyiha pul oqimi (moliyalashtirishsiz)")
line("capex", "Kapital xarajatlar", lambda i, c: f"=-({'%s' % CAPEX1 if i == 0 else 0}+{CAPEX2 if i == 1 else 0})")
line("dwc", "Aylanma kapital o'zgarishi",
     lambda i, c: f"=-{c}{R['rev']}*{K['wc']}" if i == 0 else f"=-({c}{R['rev']}-{L(2+i-1)}{R['rev']})*{K['wc']}")
line("utax", "Soliq (moliyalashtirishsiz)", lambda i, c: f"=-MAX(0,{c}{R['ebit']})*{K['tax']}")
line("fcf", "ERKIN PUL OQIMI (FCF)", lambda i, c: f"={c}{R['ebitda']}+{c}{R['capex']}+{c}{R['dwc']}+{c}{R['utax']}",
     bold=True, total=True)
line("cfcf", "Kumulyativ FCF", lambda i, c: f"={c}{R['fcf']}" if i == 0 else f"={L(2+i-1)}{R['cfcf']}+{c}{R['fcf']}")
line("dfcf", "Diskontlangan FCF", lambda i, c: f"={c}{R['fcf']}/(1+{K['disc']})^{i}")

sec("Moliyalashtirish")
line("draw", "Kredit olish", lambda i, c: f"=-{c}{R['capex']}*{K['loan_sh']}")
line("repay", "Kredit asosiy qarzini qaytarish",
     lambda i, c: f"=IF(AND({i}>=2,{i}<=6),SUM(${cols[0]}${R['draw']}:${cols[-1]}${R['draw']})/5,0)",
     note="1,5 yil imtiyozli davr, 5 yilda teng qismlarda")
line("bal", "Kredit qoldig'i (yil oxiri)",
     lambda i, c: f"={c}{R['draw']}-{c}{R['repay']}" if i == 0 else f"={L(2+i-1)}{R['bal']}+{c}{R['draw']}-{c}{R['repay']}")
line("grant", "Grant", lambda i, c: f"={K['grant']}" if i == 0 else "=0")
line("res", "Aylanma mablag' uchun qo'shimcha kapital", RESERVE, inp=True,
     note="Ishga tushirish davridagi zararni va aylanma kapitalni qoplash uchun")
line("eq", "Ta'sischi / investor ulushi (jami)",
     lambda i, c: f"=MAX(0,-{c}{R['capex']}-{c}{R['draw']}-{c}{R['grant']})+{c}{R['res']}",
     note="CAPEXning kredit va grantdan tashqari qismi + qo'shimcha kapital")
line("cf_eq", "Sof pul oqimi (moliyalashtirish bilan)",
     lambda i, c: f"={c}{R['ebitda']}+{c}{R['capex']}+{c}{R['dwc']}-{c}{R['taxp']}-{c}{R['int']}+{c}{R['draw']}-{c}{R['repay']}+{c}{R['grant']}+{c}{R['eq']}",
     bold=True)
line("cash", "Kassa qoldig'i (yil oxiri)",
     lambda i, c: f"={c}{R['cf_eq']}" if i == 0 else f"={L(2+i-1)}{R['cash']}+{c}{R['cf_eq']}", bold=True, total=True)
line("dscr", "Qarzni qoplash koeffitsienti (DSCR)",
     lambda i, c: f"=IF(({c}{R['int']}+{c}{R['repay']})=0,0,{c}{R['ebitda']}/({c}{R['int']}+{c}{R['repay']}))",
     '0.00"x"', note="Banklar odatda ≥1,2x talab qiladi")
# interest formula now
for i, c in enumerate(cols):
    opening = "0" if i == 0 else f"{L(2+i-1)}{R['bal']}"
    pg[f"{c}{R['int']}"] = f"=({opening}+{c}{R['bal']})/2*{K['rate']}"

row[0] += 1
sec("Investitsiya samaradorligi")
kp = {}
for key, label, f, fmt in [
    ("tinv", "Jami investitsiya (CAPEX)", f"={CAPEXT}", USD),
    ("npv", "NPV (sof joriy qiymat, 7 yil, qoldiq qiymatsiz)", f"=SUM({cols[0]}{R['dfcf']}:{cols[-1]}{R['dfcf']})", USD),
    ("irr", "IRR (ichki daromadlilik me'yori, 7 yil)", f"=IRR({cols[0]}{R['fcf']}:{cols[-1]}{R['fcf']})", PCT),
    ("tv", "Qoldiq qiymat (7-yil FCF × 5)", f"={cols[-1]}{R['fcf']}*5/(1+{K['disc']})^7", USD),
    ("npv_tv", "NPV qoldiq qiymat bilan", None, USD),
    ("pb", "Oddiy qoplanish muddati (yil, 0-yildan)", None, '0.0'),
    ("be", "Zararsizlik nuqtasi — plyonka hajmi (t/yil, 4-yil xarajatlari bo'yicha)", None, NUM),
]:
    rr = row[0]
    kp[key] = rr
    pg.cell(row=rr, column=1, value=label).font = BOLD
    if key == "npv_tv":
        f = f"=B{kp['npv']}+B{kp['tv']}"
    if key == "pb":
        # first year where cumulative FCF >= 0, with linear interpolation
        parts = []
        for i in range(1, NY):
            c, pc = cols[i], cols[i - 1]
            parts.append(f"IF(AND({pc}{R['cfcf']}<0,{c}{R['cfcf']}>=0),{i-1}+(-{pc}{R['cfcf']})/{c}{R['fcf']},")
        f = "=" + "".join(parts) + '"7 yildan ortiq"' + ")" * (NY - 1)
    if key == "be":
        c = cols[4]
        f = (f"=({c}{R['fc']}-({c}{R['rev_pha']}-{c}{R['vc_pha']}))/"
             f"(({K['p_film']}*(1-{K['sell']}))-{c}{R['vc_film']}/{c}{R['film']})")
    cc = pg.cell(row=rr, column=2, value=f)
    cc.number_format = fmt
    cc.font = BOLD
    cc.fill = FILL_KEY if key in ("npv", "irr", "pb") else FILL_TOT
    row[0] += 1
pg.cell(row=kp["be"], column=NY + 2,
        value="Ichki PHA tannarxi va PHA sotuvi doimiy deb olingan").font = Font(name=F, size=9, color="555555")
pg.column_dimensions["A"].width = 52
for c in cols:
    pg.column_dimensions[c].width = 13
pg.column_dimensions[NOTE].width = 60
pg.freeze_panes = "B4"

# ---------------------------------------------------------------- Sezgirlik
sz = wb.create_sheet("Sezgirlik")
title(sz, "Sezgirlik tahlili — 4-yil (2031) EBITDA, $", 6)
sz["A2"] = "Satrlar: paket narxi o'zgarishi; ustunlar: PBAT narxi o'zgarishi. Boshqa barcha omillar o'zgarmas."
sz["A2"].font = Font(name=F, italic=True, size=9)
c4 = cols[4]
base_e = f"Prognoz!${c4}${R['ebitda']}"
film_t = f"Prognoz!${c4}${R['film']}"
rev_t = f"Prognoz!${c4}${R['rev']}"
pbat_t = f"{film_t}*(1-Prognoz!${c4}${R['share']}-{K['add_sh']})*{K['pbat_sh']}*{K['pbat']}*(1+{K['scrap']})"
sz["A4"] = "Paket narxi \\ PBAT narxi"
sz["A4"].font = BOLD
dps = [-0.2, -0.1, 0, 0.1, 0.2]
for j, d in enumerate(dps):
    c = sz.cell(row=4, column=2 + j, value=d)
    c.number_format = '+0%;-0%;0%'
    c.font = Font(name=F, bold=True, color="0000FF")
for i, d in enumerate(dps):
    rr = 5 + i
    c = sz.cell(row=rr, column=1, value=d)
    c.number_format = '+0%;-0%;0%'
    c.font = Font(name=F, bold=True, color="0000FF")
    for j in range(len(dps)):
        cl = L(2 + j)
        f = (f"={base_e}+$A{rr}*{film_t}*{K['p_film']}*(1-{K['sell']})-{cl}$4*{pbat_t}")
        cc = sz.cell(row=rr, column=2 + j, value=f)
        cc.number_format = USD
        if i == 2 and j == 2:
            cc.fill = FILL_KEY
sz["A11"] = "Asosiy xulosa: rentabellik eng ko'p sotish narxiga, keyin PBAT narxiga bog'liq."
sz["A11"].font = Font(name=F, italic=True)
sz.column_dimensions["A"].width = 26
for cl in "BCDEF":
    sz.column_dimensions[cl].width = 14

# ---------------------------------------------------------------- Eski va yangi
es = wb.create_sheet("Eski vs Yangi")
title(es, "Dastlabki taqdimotdagi raqamlar va tuzatilgan qiymatlar", 4)
header(es, 3, ["Ko'rsatkich", "Dastlabki taqdimot", "Tuzatilgan (model)", "Nima uchun o'zgardi"])
cmp = [
    ("1 t plastik sotish narxi", 250, f"={K['p_film']}",
     "250 $/t polietilendan (≈1 200 $/t) ham 5 baravar arzon. PHA bozor narxi 3 400–3 700 $/t (IMARC, 2025), kompost bo'ladigan paket ≈3 900 $/t."),
    ("1 t tannarx", 95, f"=Tannarx!C{ft}",
     "Faqat yog' xomashyosining o'zi ≈900 $/t PHA (1,8 t yog' kerak). Elektr, gaz, kimyoviy moddalar, PBAT/PLA qo'shilgan."),
    ("1 t xomashyo xarajati", 50, f"=Tannarx!C5", "1 t PHA uchun ~1,8 t uglerod manbai kerak; 50 $ jismonan mumkin emas."),
    ("Oylik ishlab chiqarish (t, barqaror yil)", 50, f"=Prognoz!{cols[4]}{R['film']}/12",
     "25 000 $ lik uskuna bilan oyiga 50 t PHA olish mumkin emas (≈ 600 m³ fermenter kerak). Yangi model: oyiga ~71 t tayyor paket/plyonka va ~4,8 t PHA."),
    ("Jami boshlang'ich investitsiya ($)", 60000, f"={CAPEXT}",
     "Real narxlar: 10 m³ fermenter ≈ 50–90 ming $, separator ≈ 70 ming $, ekstruder ≈ 85 ming $."),
    ("Yillik foyda ($)", 93000, f"=Prognoz!{cols[4]}{R['ni']}", "Barqaror yil (2031) sof foydasi, soliq va kredit foizi bilan."),
    ("Investitsiya qaytimi", "8 oy", f"=Prognoz!B{kp['pb']}", "Kimyo-biotexnologiya korxonasi uchun 4–6 yil — sanoat uchun odatiy."),
]
for i, (lab, old, new, why) in enumerate(cmp):
    rr = 4 + i
    es.cell(row=rr, column=1, value=lab).font = BOLD
    c = es.cell(row=rr, column=2, value=old)
    c.font = BLUE
    if isinstance(old, (int, float)):
        c.number_format = USD if "(t)" not in lab else NUM
    c2 = es.cell(row=rr, column=3, value=new)
    c2.number_format = (NUM1 if "(t)" in lab else ('0.0" yil"' if "qaytimi" in lab else USD))
    c2.font = GREEN
    es.cell(row=rr, column=4, value=why).alignment = WRAP
    es.row_dimensions[rr].height = 45
es.column_dimensions["A"].width = 34
es.column_dimensions["B"].width = 18
es.column_dimensions["C"].width = 18
es.column_dimensions["D"].width = 90

for ws in wb.worksheets:
    for row_ in ws.iter_rows():
        for c in row_:
            if c.font and c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
wb.save(OUT)
print("saved", OUT, R, kp)
