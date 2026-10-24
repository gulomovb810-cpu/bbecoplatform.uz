import json, sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1], data_only=True)
pg = wb["Prognoz"]; R = {}
for r in pg.iter_rows(min_row=4, values_only=True):
    if r[0]: R[r[0]] = list(r[1:9])
B = lambda lab: R[lab][0]
cx = wb["CAPEX"]
items = [r[1:8] for r in cx.iter_rows(min_row=4, values_only=True) if isinstance(r[0], int)]
summ = {}
for r in cx.iter_rows(min_row=4, values_only=True):
    m = {"Uskunalar":"eq","Yetkazish, bojxona rasmiylashtiruvi":"frt","Montaj, quvurlar, elektr ulanish":"inst","Boshqa (bino, sertifikat, loyihalash)":"oth","Kutilmagan xarajatlar":"cont","JAMI CAPEX":"tot"}
    if r[1] in m and r[0] is None: summ[m[r[1]]] = [r[5], r[6], r[7]]
xo = wb["Xodimlar"]
staff = [r[:4] for r in xo.iter_rows(min_row=4, max_row=14, values_only=True)]
tn = list(wb["Tannarx"].iter_rows(min_row=5, values_only=True))
pha_lines = [(r[0], r[2]) for r in tn[:7]]
fi = [i for i, r in enumerate(tn) if r[0] == "PBAT"][0]
film_lines = [(r[0], r[2]) for r in tn[fi:fi+6]]
D = dict(
  capex=summ["tot"][2], capex1=summ["tot"][0], capex2=summ["tot"][1], cx=summ,
  loan=sum(R["Kredit olish"]), grant=sum(R["Grant"]), equity=sum(R["Ta'sischi / investor ulushi (jami)"]),
  rev4=R["JAMI DAROMAD"][4], ebitda4=R["EBITDA"][4], ni4=R["SOF FOYDA"][4], film4=R["Paket/plyonka ishlab chiqarish (t)"][4],
  phasale4=R["PHA sotish (t)"][4],
  npv=B("NPV (sof joriy qiymat, 7 yil, qoldiq qiymatsiz)"), npv_tv=B("NPV qoldiq qiymat bilan"),
  irr=B("IRR (ichki daromadlilik me'yori, 7 yil)"), payback=B("Oddiy qoplanish muddati (yil, 0-yildan)"),
  be=B("Zararsizlik nuqtasi — plyonka hajmi (t/yil, 4-yil xarajatlari bo'yicha)"),
  staff=staff, staff1=sum(s[1] for s in staff), staff2=sum(s[2] for s in staff),
  pay1=xo["E16"].value, pay2=xo["F16"].value,
  pha_lines=pha_lines, film_lines=film_lines, pha_vc=tn[7][2],
  film_vc=[r[2] for r in tn if r[0]=="Paket/plyonka o'zgaruvchan tannarxi"][0],
  film_margin=[r[2] for r in tn if r[0] and r[0].startswith("Marjinal")][0],
  years=[str(2027+i) for i in range(8)],
  film=R["Paket/plyonka ishlab chiqarish (t)"], pha=R["PHA ishlab chiqarish (t)"], rev=R["JAMI DAROMAD"], gp=R["MARJINAL FOYDA"],
  fc=R["Jami doimiy xarajat"], ebitda=R["EBITDA"], dep=R["Amortizatsiya"], int=R["Kredit foizi"], ni=R["SOF FOYDA"],
  fcf=R["ERKIN PUL OQIMI (FCF)"], cfcf=R["Kumulyativ FCF"], dscr=R["Qarzni qoplash koeffitsienti (DSCR)"][2:7],
  sens=[[c for c in r[1:6]] for r in wb["Sezgirlik"].iter_rows(min_row=5, max_row=9, values_only=True)],
  capex_items=items,
)
D["funding"] = D["loan"] + D["grant"] + D["equity"]
json.dump(D, open(sys.argv[2], "w"), ensure_ascii=False, default=float, indent=0)
print({k: D[k] for k in ["capex","loan","grant","equity","funding","payback","irr","staff1","staff2","pay1","pay2","pha_vc","film_vc"]})
